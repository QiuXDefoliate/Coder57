from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "交付材料"
OUT.mkdir(exist_ok=True)

NAVY = "173E36"
PALE = "EAF2EF"
GRAY = "F3F4F2"
BORDER = "D9D9D9"

def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)

def set_cell_border(cell, color=BORDER):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), color)

def set_run_font(run, name="Microsoft YaHei", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)

def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn("w:pBdr"))
    if border is not None:
        p_pr.remove(border)

def base_doc(title, subtitle):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    for style_name, size, bold in [("Normal", 10.5, False), ("Title", 24, True), ("Heading 1", 17, True), ("Heading 2", 13, True)]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        remove_paragraph_border(style)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Normal"].paragraph_format.line_spacing = 1.25
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    p = doc.add_paragraph(style="Title")
    remove_paragraph_border(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(title), size=24, bold=True)
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p2.add_run(subtitle), size=11, color="555555")
    doc.add_paragraph()
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(intro.add_run("澄心银行 AI Agent Web MVP"), size=14, bold=True, color=NAVY)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(meta.add_run("版本 1.0    2026 年 9 月"), size=10, color="666666")
    doc.add_page_break()
    return doc

def add_p(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        set_run_font(p.add_run(bold_lead), bold=True)
        set_run_font(p.add_run(text[len(bold_lead):]))
    else:
        set_run_font(p.add_run(text))
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run_font(p.add_run(item))

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = True
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, NAVY)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(h), size=9.5, bold=True, color="FFFFFF")
    for r_index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            if r_index % 2:
                set_cell_fill(cells[i], GRAY)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), size=9.2)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph()
    return table

def add_footer(doc, label):
    for section in doc.sections:
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(p.add_run(label), size=8, color="777777")

def build_technical():
    doc = base_doc("银行 AI Agent 技术文档", "系统架构 核心算法与安全设计")
    doc.add_heading("1 项目概述", level=1)
    add_p(doc, "本项目实现一套面向个人用户的银行 AI Agent 演示系统。系统提供自然语言转账、定时转账、AA 收款、账单分析、理财产品比较、卡片管理、订阅管理和生日计划。第一版只处理模拟数据，不连接真实银行、证券、征信、日历或商户。")
    add_p(doc, "核心结论：模型负责理解和解释，后端负责资金校验、确认状态、强认证、幂等执行和审计。任何资金动作都不能由模型直接修改账户余额。", "核心结论：")

    doc.add_heading("2 系统架构", level=1)
    add_table(doc, ["层级", "技术", "职责"], [
        ["Web 前端", "Next.js 15 React 19", "聊天工作台、确认卡、账单、理财、卡片、订阅和审计页面"],
        ["业务 API", "NestJS 11", "意图路由、参数校验、操作状态机、模拟银行工具接口"],
        ["模型服务", "DeepSeek OpenAI 兼容接口", "意图分类和自然语言理解；不可直接执行交易"],
        ["任务系统", "BullMQ Redis", "定时转账、续费提醒、生日计划和账单报告任务"],
        ["数据层", "Prisma PostgreSQL", "账户、交易、操作、审计事件的数据模型"],
    ], [1.1, 2.0, 3.5])
    add_p(doc, "演示模式默认使用进程内种子数据，便于评审环境直接运行。项目已经提供 PostgreSQL Schema、Redis worker 和 Docker Compose，后续可把 DemoStore 替换为持久化仓储。")

    doc.add_heading("3 Agent 处理流程", level=1)
    add_table(doc, ["阶段", "输入", "输出与约束"], [
        ["理解", "用户自然语言或语音转写", "识别转账、账单、理财、卡片、订阅或自动化意图"],
        ["准备", "结构化参数", "校验收款人、余额、限额和产品适配，生成待确认操作"],
        ["确认", "用户核对确认卡", "状态改为等待强认证，不扣款"],
        ["执行", "6 位演示认证码", "后端再次校验并按幂等键执行"],
        ["审计", "操作结果", "记录准备、确认、执行或拒绝事件"],
    ], [0.9, 2.4, 3.3])
    add_p(doc, "写操作统一携带 operationId、idempotencyKey、riskLevel、status、expiresAt 和业务详情。确认卡十分钟后过期，重复执行同一幂等键不会重复扣款。")

    doc.add_heading("4 核心功能", level=1)
    add_bullets(doc, [
        "智能转账：支持姓名、已绑定手机号、语音转写和定时执行；重名联系人必须消歧。",
        "AA 收款：识别总金额、参与人和是否排除发起人，生成模拟收款请求。",
        "账单分析：按商户与描述分类，展示分类占比，并由确定性规则标记异常交易。",
        "理财操作：先检查风险等级，再展示适配产品；申购必须经过确认和强认证。",
        "卡片管理：展示脱敏卡号和交易限制；挂失执行后更新卡片状态。",
        "订阅管理：识别周期扣款；只有标记为 direct_cancel 的商户才能发起银行侧取消。",
        "跨场景联动：识别生日后先准备资金预留，鲜花和蛋糕订单需要再次单独确认。",
    ])

    doc.add_heading("5 核心算法说明", level=1)
    doc.add_heading("5.1 意图识别与回退", level=2)
    add_p(doc, "配置 DeepSeek API Key 时，系统调用 deepseek-flash 输出结构化意图。请求关闭思考模式并要求 JSON 输出。模型调用失败或未配置 Key 时，本地安全解析器根据关键词继续提供演示功能。模型返回的意图仅用于路由，后端仍重新校验业务参数。")
    doc.add_heading("5.2 金额与收款人解析", level=2)
    add_p(doc, "金额解析优先识别“转 500 元”“预留 1000 元”等与动作相邻的金额，避免把 11 位手机号误判为金额。收款人先按手机号精确匹配，再按姓名匹配；匹配到多位同名联系人时停止流程并要求用户指定银行尾号。")
    doc.add_heading("5.3 异常交易识别", level=2)
    add_p(doc, "演示规则包括凌晨大额、非常用城市、金额显著高于个人基线、短时重复和陌生商户。规则引擎负责生成信号，AI 只解释信号，不直接判断欺诈，也不自动冻结卡片。")
    doc.add_heading("5.4 幂等执行", level=2)
    add_p(doc, "服务端在执行前查询 idempotencyKey。已执行的键直接返回原操作结果，不再次修改余额。执行前同时检查状态必须为 AWAITING_AUTH、确认未过期、认证码正确且可用余额足够。")

    doc.add_heading("6 提示词设计", level=1)
    add_p(doc, "提示词包位于 packages/prompts/src/index.ts，版本号为 bank-agent-cn-1.0.0。它包含总控、意图路由、转账提取、账单分类、异常解释、报告生成、理财适当性、卡片与订阅、跨场景规划、安全审查十类提示词。")
    add_table(doc, ["原则", "实现"], [
        ["最小权限", "模型只能建议工具调用，不能直接修改余额或状态"],
        ["数据可信", "金额、余额和产品条款只引用后端工具结果"],
        ["隐私保护", "卡号、身份证号、密码和验证码在模型调用前拦截或移除"],
        ["防提示注入", "商户名、备注和工具结果只作为数据，不作为系统指令"],
        ["风险披露", "禁止保本、稳赚和确定收益表述"],
    ], [1.3, 5.3])

    doc.add_heading("7 安全设计", level=1)
    add_bullets(doc, [
        "资金、高风险卡片操作和订阅取消使用分级确认。",
        "完整卡号、身份证号、密码和验证码禁止进入模型上下文。",
        "DeepSeek user_id 使用不可逆内部标识，不包含手机号或姓名。",
        "确认卡显示付款账户尾号、收款人、脱敏手机号、金额、手续费和执行时间。",
        "审计日志记录操作编号、动作、结果和脱敏详情。",
        "模型服务不可用时，账户查询和已创建任务不依赖模型继续运行。",
    ])

    doc.add_heading("8 部署与运行", level=1)
    add_table(doc, ["步骤", "命令"], [
        ["准备配置", "Copy-Item .env.example .env"],
        ["安装依赖", "npm install"],
        ["启动演示", "npm run dev"],
        ["启动数据库与 Redis", "docker compose up -d"],
        ["生成 Prisma Client", "npm run db:generate"],
        ["启动 worker", "npm run dev:worker"],
    ], [1.6, 5.0])
    add_p(doc, "浏览器访问 http://localhost:3000。演示强认证码为 123456。若需启用 DeepSeek，在根目录 .env 中填写 DEEPSEEK_API_KEY。")

    doc.add_heading("9 测试结果与限制", level=1)
    add_bullets(doc, [
        "6 个 API 单元测试通过，覆盖重名消歧、确认前不扣款、认证执行、幂等、提示注入、挂失和订阅取消边界。",
        "API、Web 和 worker 通过 TypeScript 类型检查与生产构建。",
        "HTTP 冒烟测试确认转账执行前后余额差额为 500 元，并生成准备、确认、执行三条审计事件。",
        "当前不接真实银行接口，不具备生产级 KYC、短信、生物识别、监管报送或真实商户履约能力。",
    ])
    add_footer(doc, "银行 AI Agent 技术文档")
    path = OUT / "01_银行AI_Agent技术文档.docx"
    doc.save(path)
    return path

def build_security():
    doc = base_doc("银行 AI Agent 安全自评报告", "权限分级 风险识别与整改清单")
    doc.add_heading("1 自评范围", level=1)
    add_p(doc, "本报告评估澄心银行 AI Agent Web MVP 的演示代码、模型接入、资金操作状态机、日志和基础部署配置。报告结论仅适用于当前模拟系统，不代表通过银行业等保、渗透测试、代码审计或监管验收。")

    doc.add_heading("2 数据与资产", level=1)
    add_table(doc, ["资产", "敏感级别", "当前处理方式"], [
        ["账户余额与交易", "高", "保存在演示存储；展示时限制为当前演示用户"],
        ["卡片信息", "高", "只展示尾号；完整卡号不得进入模型"],
        ["联系人手机号", "中", "内部匹配，向用户和模型输出脱敏号码"],
        ["认证信息", "高", "演示码仅用于状态机演示，不写入日志"],
        ["模型会话", "中", "调用前进行敏感字段检测和移除"],
        ["审计记录", "中", "记录操作与结果，不记录密码或验证码"],
    ], [1.5, 1.1, 4.0])

    doc.add_heading("3 风险评估", level=1)
    add_table(doc, ["风险", "等级", "已有控制", "剩余问题"], [
        ["模型越权执行资金动作", "高", "模型无余额写权限，写操作必须经过确认和认证", "生产环境需独立授权服务"],
        ["提示词注入绕过规则", "高", "检测跳过确认、伪造回执和系统提示词请求", "需要持续红队测试"],
        ["敏感数据发送给模型", "高", "正则检测卡号、身份证号、密码和验证码", "需增加数据分类网关"],
        ["重复请求导致重复扣款", "高", "operationId 与 idempotencyKey", "需在数据库唯一约束下验证并发"],
        ["异常交易误报", "中", "规则只给出信号，AI 不断言欺诈", "需基于真实样本调参"],
        ["订阅取消状态虚假", "中", "按商户 capability 决定直接取消或给出步骤", "真实商户接口尚未接入"],
        ["依赖服务不可用", "中", "DeepSeek 失败时回退本地解析", "需补充监控和告警"],
    ], [1.6, 0.7, 2.5, 2.0])

    doc.add_heading("4 权限与交易安全", level=1)
    add_bullets(doc, [
        "R0 查询操作可直接执行，R1 低风险写操作需要确认，R2 和 R3 资金或高风险操作需要确认与强认证。",
        "转账、申购、挂失、订阅取消、资金预留和商户下单统一生成 PendingOperation。",
        "服务端拒绝未确认、确认过期、认证失败、余额不足和状态不正确的请求。",
        "执行成功后写入审计事件；同一幂等键再次调用返回原结果。",
    ])

    doc.add_heading("5 模型安全", level=1)
    add_bullets(doc, [
        "系统提示词明确禁止直接改变资金、泄露数据、跳过确认和虚构成功回执。",
        "交易备注、商户名、网页文字和工具结果均视为数据，不视为指令。",
        "模型产生的金额、产品或工具参数不直接可信，后端必须重新校验。",
        "理财输出必须先读取风险等级，并禁止保本、稳赚或确定收益描述。",
    ])

    doc.add_heading("6 自查清单", level=1)
    add_table(doc, ["检查项", "结果", "说明"], [
        ["敏感操作有确认步骤", "通过", "统一 prepare confirm execute"],
        ["确认前不修改余额", "通过", "单元测试与 HTTP 冒烟测试覆盖"],
        ["重复执行不会重复扣款", "通过", "演示存储维护已执行幂等键"],
        ["卡号和手机号脱敏", "部分通过", "展示层已脱敏，仍需生产级数据网关"],
        ["生产身份认证", "未完成", "当前使用固定 6 位演示码"],
        ["数据库并发与灾备", "未完成", "当前默认使用内存数据"],
        ["渗透测试与依赖扫描", "未完成", "交付前需由安全团队执行"],
        ["监管与隐私合规评审", "未完成", "需要真实业务主体和数据流程后开展"],
    ], [2.2, 1.0, 3.4])

    doc.add_heading("7 结论与整改优先级", level=1)
    add_p(doc, "当前 MVP 满足演示环境的基本安全边界：模型不能直接动账，资金操作有确认和强认证，重复提交具有幂等保护，敏感输入会被拦截。系统尚不满足真实银行生产要求。")
    add_table(doc, ["优先级", "整改项"], [
        ["P0", "接入真实身份认证、授权服务和数据库事务，移除固定演示认证码"],
        ["P0", "建立敏感数据分类与模型调用网关，执行内容审计和最小化传输"],
        ["P1", "补充限流、会话保护、并发幂等测试和操作告警"],
        ["P1", "开展依赖扫描、静态代码审计、渗透测试和提示词红队测试"],
        ["P2", "建立异常模型评估集、误报处理和人工复核流程"],
    ], [1.0, 5.6])
    add_footer(doc, "银行 AI Agent 安全自评报告")
    path = OUT / "04_银行AI_Agent安全自评报告.docx"
    doc.save(path)
    return path

if __name__ == "__main__":
    for output in (build_technical(), build_security()):
        print(output)
